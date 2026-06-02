from __future__ import annotations

import ast
import json
import math
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Callable, Sequence

import pandas as pd
from openpyxl import load_workbook

from scrapper.source_rules import get_source_rules
from scrapper.storage import JobStore
from scrapper.text import clean_text, contains_any_terms, has_excluded_contract_signal, url_key


MONEY_RE = re.compile(r"(\d[\d\s\u202f.,]*)")
ISO_DATE_PREFIX_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")
EU_DATE_RE = re.compile(r"^(?P<day>\d{1,2})/(?P<month>\d{1,2})/(?P<year>\d{4})$")
NATURAL_LANGUAGE_DATE_RE = re.compile(
    r"(?:Publié le\s+)?(?P<day>\d{1,2})\s+(?P<month>[A-Za-zÀ-ÿ]+)(?:\s+(?P<year>\d{4}))?"
)
RELATIVE_DATE_RE = re.compile(
    r"il\s+y\s+a\s+(?P<count>\d+)\s*(?P<unit>min|mn|minute?s?|h|heure?s?|j|jour?s?)",
    re.IGNORECASE,
)
HYPERLINK_RE = re.compile(r'^=HYPERLINK\("(?P<url>[^"]+)"(?:,"[^"]*")?\)$', re.IGNORECASE)
LINK_COLUMNS = ("Lien", "Postuler", "Offre source")
STATUS_OPTIONS = ("En Attente", "Postulé", "Refusé")
DEFAULT_STATUS = STATUS_OPTIONS[0]
MAX_EXCEL_ROW_INDEX = 1_048_575
FRENCH_MONTHS = {
    "janvier": 1,
    "fevrier": 2,
    "février": 2,
    "mars": 3,
    "avril": 4,
    "mai": 5,
    "juin": 6,
    "juillet": 7,
    "aout": 8,
    "août": 8,
    "septembre": 9,
    "octobre": 10,
    "novembre": 11,
    "decembre": 12,
    "décembre": 12,
}
LOCALIZED_TRACKER_COLUMNS = (
    "Postulé",
    "Rejeté",
    "Lien",
    "Statut",
    "Postuler",
    "Offre source",
    "Autres liens",
    "Intitulé",
    "Entreprise",
    "Lieu",
    "Contrat",
    "Télétravail",
    "Salaire min (€ annuel)",
    "Salaire max (€ annuel)",
    "Salaire",
    "Publié le",
    "Sources",
    "Source principale",
    "Doublons",
    "Description",
    "URL entreprise",
    "Email",
    "Recherche",
    "Zone recherchée",
    "Score complétude",
    "Vu le",
    "Dernière vue",
)
DEFAULT_LOCALIZED_TRACKER_COLUMNS = (
    "Postulé",
    "Rejeté",
    "Lien",
    "Intitulé",
    "Entreprise",
    "Lieu",
    "Contrat",
    "Télétravail",
    "Salaire min (€ annuel)",
    "Salaire max (€ annuel)",
    "Salaire",
    "Publié le",
    "Sources",
    "Recherche",
)


def export_jobs(
    store: JobStore,
    export_dir: Path,
    excel_columns: Sequence[str] | None = None,
    repair_fields_enabled_by_default: bool = True,
    repair_fields_by_source: dict[str, bool] | None = None,
    include_terms: tuple[str, ...] = (),
    exclude_terms: tuple[str, ...] = (),
    location_allowed: Callable[[object | None], bool] | None = None,
) -> tuple[Path, Path, Path, int]:
    export_dir.mkdir(parents=True, exist_ok=True)
    csv_path = export_dir / "jobs.csv"
    raw_csv_path = export_dir / "jobs_raw.csv"
    excel_path = export_dir / "jobs.xlsx"

    canonical = _excel_safe(store.canonical_dataframe())
    raw = _excel_safe(store.raw_dataframe())
    localized = _select_tracker_columns(
        _localized_job_tracker(
            canonical,
            repair_fields_enabled_by_default=repair_fields_enabled_by_default,
            repair_fields_by_source=repair_fields_by_source,
        ),
        excel_columns or DEFAULT_LOCALIZED_TRACKER_COLUMNS,
    )
    localized = _merge_existing_tracker_rows(
        localized,
        excel_path,
        csv_path,
        include_terms,
        exclude_terms,
        location_allowed,
    )
    localized = _sort_tracker_for_export(localized)

    localized.to_csv(csv_path, index=False)
    raw.to_csv(raw_csv_path, index=False)

    with pd.ExcelWriter(excel_path, engine="xlsxwriter", date_format="dd/mm/yyyy", datetime_format="dd/mm/yyyy") as writer:
        localized.to_excel(writer, index=False, sheet_name="Offres")
        raw.drop(columns=["raw_json"], errors="ignore").to_excel(writer, index=False, sheet_name="Données brutes")
        _format_jobs_sheet(writer, "Offres", localized)
        _format_raw_sheet(writer, "Données brutes", raw.drop(columns=["raw_json"], errors="ignore"))

    return excel_path, csv_path, raw_csv_path, len(localized)


def _merge_existing_tracker_rows(
    current: pd.DataFrame,
    excel_path: Path,
    csv_path: Path,
    include_terms: tuple[str, ...] = (),
    exclude_terms: tuple[str, ...] = (),
    location_allowed: Callable[[object | None], bool] | None = None,
) -> pd.DataFrame:
    previous = _read_existing_tracker(excel_path, csv_path)
    if previous.empty:
        return current

    current = current.copy()
    previous = previous[
        previous.apply(
            _tracker_row_allowed,
            axis=1,
            include_terms=include_terms,
            exclude_terms=exclude_terms,
            location_allowed=location_allowed,
        )
    ]
    previous = previous.apply(_normalize_previous_tracker_row, axis=1)
    previous_by_key: dict[str, pd.Series] = {}
    for _, row in previous.iterrows():
        key = _tracker_row_key(row)
        if key and key not in previous_by_key:
            previous_by_key[key] = row

    current_keys: set[str] = set()
    for index, row in current.iterrows():
        key = _tracker_row_key(row)
        if not key:
            continue
        current_keys.add(key)
        previous_row = previous_by_key.get(key)
        if previous_row is not None:
            _preserve_tracker_state(current, index, previous_row)

    retained_previous = [
        row
        for _, row in previous.iterrows()
        if (key := _tracker_row_key(row)) and key not in current_keys
    ]
    if not retained_previous:
        if "Rejeté" in current.columns:
            current["Rejeté"] = current["Rejeté"].apply(_coerce_checkbox_value)
        return current
    retained = pd.DataFrame(retained_previous).reindex(columns=current.columns)
    merged = pd.concat([current, retained], ignore_index=True)
    if "Rejeté" in merged.columns:
        merged["Rejeté"] = merged["Rejeté"].apply(_coerce_checkbox_value)
    return merged


def _sort_tracker_for_export(frame: pd.DataFrame) -> pd.DataFrame:
    if frame.empty or "Publié le" not in frame.columns:
        return frame
    sorted_frame = frame.copy()
    sorted_frame["__published_sort"] = pd.to_datetime(
        sorted_frame["Publié le"].map(_date_value),
        errors="coerce",
    )
    sorted_frame = sorted_frame.sort_values(
        by="__published_sort",
        ascending=False,
        na_position="last",
        kind="stable",
    )
    return sorted_frame.drop(columns=["__published_sort"]).reset_index(drop=True)


def _tracker_row_allowed(
    row: pd.Series,
    include_terms: tuple[str, ...],
    exclude_terms: tuple[str, ...],
    location_allowed: Callable[[object | None], bool] | None = None,
) -> bool:
    title = row.get("Intitulé")
    if exclude_terms and has_excluded_contract_signal(title, row.get("Contrat"), "", exclude_terms):
        return False
    if include_terms and not contains_any_terms((title,), include_terms):
        return False
    return location_allowed is None or location_allowed(row.get("Lieu"))


def _normalize_previous_tracker_row(row: pd.Series) -> pd.Series:
    normalized = row.copy()
    if "Statut" in normalized.index:
        normalized["Statut"] = _normalize_status(normalized.get("Statut"))
    if "Postulé" in normalized.index:
        normalized["Postulé"] = _checkbox_value_from_previous(normalized)
    if "Rejeté" in normalized.index:
        normalized["Rejeté"] = _coerce_checkbox_value(normalized.get("Rejeté"))
    if "Lieu" in normalized.index:
        normalized["Lieu"] = _town_fr(normalized.get("Lieu"))
    if "Publié le" in normalized.index:
        normalized["Publié le"] = _date_value(normalized.get("Publié le"))
    for column in ("Lien", "Postuler", "Offre source", "URL entreprise"):
        if column not in normalized.index:
            continue
        url = _extract_hyperlink_url(normalized.get(column))
        if url:
            normalized[column] = _hyperlink(url)
    return normalized


def _read_existing_tracker(excel_path: Path, csv_path: Path) -> pd.DataFrame:
    if excel_path.exists():
        try:
            return _read_existing_excel_tracker(excel_path)
        except Exception:  # noqa: BLE001 - a corrupt/open workbook should not block a fresh export.
            pass
    if csv_path.exists():
        try:
            return pd.read_csv(csv_path, dtype=object).fillna("")
        except Exception:  # noqa: BLE001 - fall back to a fresh export.
            pass
    return pd.DataFrame()


def _read_existing_excel_tracker(path: Path) -> pd.DataFrame:
    workbook = load_workbook(path, data_only=False, read_only=True)
    try:
        if "Offres" not in workbook.sheetnames:
            return pd.DataFrame()
        worksheet = workbook["Offres"]
        rows = worksheet.iter_rows(values_only=True)
        headers = [str(value).strip() if value is not None else "" for value in next(rows, ())]
        headers = [header for header in headers if header]
        if not headers:
            return pd.DataFrame()
        records = []
        for values in rows:
            if not values or all(value is None for value in values):
                continue
            records.append({header: values[index] if index < len(values) else "" for index, header in enumerate(headers)})
        return pd.DataFrame.from_records(records).fillna("")
    finally:
        workbook.close()


def _preserve_tracker_state(current: pd.DataFrame, index: int, previous: pd.Series) -> None:
    if "Statut" in current.columns:
        current.at[index, "Statut"] = _normalize_status(previous.get("Statut"))
    if "Postulé" in current.columns:
        current.at[index, "Postulé"] = _checkbox_value_from_previous(previous)
    if "Rejeté" in current.columns:
        current.at[index, "Rejeté"] = _coerce_checkbox_value(previous.get("Rejeté"))


def _tracker_row_key(row: pd.Series) -> str | None:
    for column in LINK_COLUMNS:
        if column not in row.index:
            continue
        extracted_url = _extract_hyperlink_url(row.get(column))
        key = url_key(extracted_url) if extracted_url else None
        if key:
            return f"url:{key}"

    title = _normalized_tracker_text(row.get("Intitulé"))
    company = _normalized_tracker_text(row.get("Entreprise"))
    location = _normalized_tracker_text(row.get("Lieu"))
    if title and company:
        return f"text:{title}|{company}|{location}"
    return None


def _extract_hyperlink_url(value: Any) -> str | None:
    text = str(value or "").strip()
    if not text:
        return None
    match = HYPERLINK_RE.match(text)
    return match.group("url") if match else text if text.startswith(("http://", "https://")) else None


def _normalized_tracker_text(value: Any) -> str:
    return clean_text(value).casefold()


def _is_blank(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return str(value).strip() == ""


def _checkbox_value_from_previous(previous: pd.Series) -> bool:
    status = str(previous.get("Statut") or "").strip()
    if status == "Postulé":
        return True
    return _coerce_checkbox_value(previous.get("Postulé"))


def _coerce_checkbox_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, float) and math.isnan(value):
        return False
    text = str(value).strip().casefold()
    return text in {"true", "1", "yes", "oui", "x", "checked", "postulé", "postule"}


def _normalize_status(value: Any) -> str:
    text = str(value or "").strip()
    folded = text.casefold()
    if folded in {"à étudier", "a etudier", "en attente", "attente", ""}:
        return DEFAULT_STATUS
    if folded in {"postulé", "postule"}:
        return "Postulé"
    if folded in {"refusé", "refuse", "rejeté", "rejete"}:
        return "Refusé"
    return text if text in STATUS_OPTIONS else DEFAULT_STATUS


def _localized_job_tracker(
    frame: pd.DataFrame,
    repair_fields_enabled_by_default: bool = True,
    repair_fields_by_source: dict[str, bool] | None = None,
) -> pd.DataFrame:
    records: list[dict[str, Any]] = []
    repair_fields_by_source = repair_fields_by_source or {}
    for _, row in frame.iterrows():
        salary_min, salary_max = _salary_bounds(row)
        best_source = str(row.get("best_source") or "").strip()
        should_repair = repair_fields_by_source.get(best_source, repair_fields_enabled_by_default)
        source_rules = get_source_rules(best_source)
        if should_repair:
            title, company, location = source_rules.repair_tracker_fields(
                row.get("title"),
                row.get("company"),
                row.get("location"),
                row.get("description"),
            )
        else:
            title = clean_text(row.get("title"))
            company = clean_text(row.get("company"))
            location = clean_text(row.get("location"))
        records.append(
            {
                "Postulé": False,
                "Rejeté": False,
                "Lien": _hyperlink(row.get("best_direct_url") or row.get("best_source_url")),
                "Statut": DEFAULT_STATUS,
                "Postuler": _hyperlink(row.get("best_direct_url") or row.get("best_source_url")),
                "Offre source": _hyperlink(row.get("best_source_url")),
                "Autres liens": _compact_links(row.get("all_source_links"), row.get("all_direct_links")),
                "Intitulé": title,
                "Entreprise": company,
                "Lieu": _town_fr(location),
                "Contrat": _contract_fr(row.get("employment_type")),
                "Télétravail": _remote_fr(row.get("is_remote")),
                "Salaire min (€ annuel)": salary_min,
                "Salaire max (€ annuel)": salary_max,
                "Salaire": _salary_label(salary_min, salary_max),
                "Publié le": _date_value(row.get("date_posted")),
                "Sources": row.get("sources"),
                "Source principale": row.get("best_source"),
                "Doublons": row.get("duplicate_count"),
                "Description": row.get("description"),
                "URL entreprise": _hyperlink(row.get("company_url")),
                "Email": row.get("emails"),
                "Recherche": clean_text(row.get("search_query")),
                "Zone recherchée": row.get("search_location"),
                "Score complétude": row.get("completeness_score"),
                "Vu le": _date_value(row.get("first_seen_at")),
                "Dernière vue": _date_value(row.get("last_seen_at")),
            }
        )
    return pd.DataFrame.from_records(records, columns=LOCALIZED_TRACKER_COLUMNS)


def _select_tracker_columns(frame: pd.DataFrame, columns: Sequence[str]) -> pd.DataFrame:
    if not columns:
        raise ValueError("At least one Excel column must be configured.")
    unknown = tuple(column for column in columns if column not in LOCALIZED_TRACKER_COLUMNS)
    if unknown:
        available = ", ".join(LOCALIZED_TRACKER_COLUMNS)
        requested = ", ".join(unknown)
        raise ValueError(f"Unknown Excel column(s): {requested}. Available columns: {available}")
    return frame.loc[:, list(columns)]


def _format_jobs_sheet(writer: pd.ExcelWriter, sheet_name: str, frame: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    workbook = writer.book
    worksheet.freeze_panes(1, 0)
    worksheet.hide_gridlines(2)

    header_format = workbook.add_format(
        {
            "bold": True,
            "font_color": "#FFFFFF",
            "bg_color": "#1F4E78",
            "align": "center",
            "valign": "vcenter",
            "text_wrap": True,
        }
    )
    checkbox_format = workbook.add_format({"align": "center", "valign": "vcenter"})
    description_format = workbook.add_format({"text_wrap": True, "valign": "top"})
    salary_format = workbook.add_format({"num_format": '# ##0" €"'})
    date_format = workbook.add_format({"num_format": "dd/mm/yyyy"})
    worksheet.set_row(0, 28, header_format)
    for index, column in enumerate(frame.columns):
        worksheet.write(0, index, column, header_format)

    checkbox_col = _optional_column_index(frame, "Postulé")
    rejected_checkbox_col = _optional_column_index(frame, "Rejeté")
    salary_min_col = _optional_column_index(frame, "Salaire min (€ annuel)")
    salary_max_col = _optional_column_index(frame, "Salaire max (€ annuel)")
    description_col = _optional_column_index(frame, "Description")
    status_col = _optional_column_index(frame, "Statut")

    for row in range(1, len(frame) + 1):
        if checkbox_col:
            value = _coerce_checkbox_value(frame.iloc[row - 1, checkbox_col - 1])
            worksheet.insert_checkbox(row, checkbox_col - 1, value, checkbox_format)
        if rejected_checkbox_col:
            value = _coerce_checkbox_value(frame.iloc[row - 1, rejected_checkbox_col - 1])
            worksheet.insert_checkbox(row, rejected_checkbox_col - 1, value, checkbox_format)
    if status_col:
        worksheet.data_validation(
            1,
            status_col - 1,
            MAX_EXCEL_ROW_INDEX,
            status_col - 1,
            {
                "validate": "list",
                "source": list(STATUS_OPTIONS),
                "input_title": "Statut",
                "input_message": "Choisir un statut.",
                "error_title": "Statut invalide",
                "error_message": "Choisissez En Attente, Postulé ou Refusé.",
            },
        )

    widths = {
        "Postulé": 10,
        "Rejeté": 10,
        "Lien": 72,
        "Statut": 15,
        "Postuler": 12,
        "Offre source": 14,
        "Autres liens": 30,
        "Intitulé": 44,
        "Entreprise": 28,
        "Lieu": 30,
        "Contrat": 18,
        "Télétravail": 14,
        "Salaire min (€ annuel)": 18,
        "Salaire max (€ annuel)": 18,
        "Salaire": 18,
        "Publié le": 14,
        "Sources": 24,
        "Description": 80,
    }
    column_formats = {}
    if checkbox_col:
        column_formats["Postulé"] = checkbox_format
    if rejected_checkbox_col:
        column_formats["Rejeté"] = checkbox_format
    if salary_min_col:
        column_formats["Salaire min (€ annuel)"] = salary_format
    if salary_max_col:
        column_formats["Salaire max (€ annuel)"] = salary_format
    if description_col:
        column_formats["Description"] = description_format
    for column in ("Publié le", "Vu le", "Dernière vue"):
        if column in frame.columns:
            column_formats[column] = date_format
    _apply_widths(worksheet, frame, widths, column_formats)
    _add_table(worksheet, "TableOffres", frame)


def _format_raw_sheet(writer: pd.ExcelWriter, sheet_name: str, frame: pd.DataFrame) -> None:
    worksheet = writer.sheets[sheet_name]
    workbook = writer.book
    header_format = workbook.add_format({"bold": True, "font_color": "#FFFFFF", "bg_color": "#595959"})
    worksheet.freeze_panes(1, 0)
    worksheet.set_row(0, None, header_format)
    for index, column in enumerate(frame.columns):
        worksheet.write(0, index, column, header_format)
    _autosize(worksheet, frame)
    _add_table(worksheet, "TableDonneesBrutes", frame)


def _add_table(worksheet, name: str, frame: pd.DataFrame) -> None:
    row_count = len(frame)
    column_count = len(frame.columns)
    if row_count < 1 or column_count < 1:
        return
    worksheet.add_table(
        0,
        0,
        row_count,
        column_count - 1,
        {
            "name": name,
            "style": "Table Style Medium 2",
            "first_column": False,
            "last_column": False,
            "banded_rows": True,
            "banded_columns": False,
            "autofilter": True,
            "columns": [{"header": column} for column in frame.columns],
        },
    )


def _salary_bounds(row: pd.Series) -> tuple[int | None, int | None]:
    minimum = _number(row.get("min_amount"))
    maximum = _number(row.get("max_amount"))
    if minimum or maximum:
        return _round_salary(minimum), _round_salary(maximum or minimum)

    salary_payload = _parse_salary_payload(row.get("salary"))
    if salary_payload:
        minimum, maximum = _bounds_from_monetary_amount(salary_payload)
        if minimum or maximum:
            return _round_salary(minimum), _round_salary(maximum or minimum)

    raw_payload = _load_json(row.get("raw_json"))
    for payload in _salary_payloads_from_raw(raw_payload):
        minimum, maximum = _bounds_from_monetary_amount(payload)
        if minimum or maximum:
            return _round_salary(minimum), _round_salary(maximum or minimum)
    return None, None


def _bounds_from_monetary_amount(payload: dict[str, Any]) -> tuple[float | None, float | None]:
    value = payload.get("value") if isinstance(payload, dict) else None
    if isinstance(value, dict):
        minimum = _number(value.get("minValue"))
        maximum = _number(value.get("maxValue"))
        fixed = _number(value.get("value"))
        unit = str(value.get("unitText") or "").upper()
        if fixed and not minimum and not maximum:
            minimum = maximum = fixed
        if unit == "MONTH":
            minimum = minimum * 12 if minimum else None
            maximum = maximum * 12 if maximum else None
        elif unit and unit not in {"YEAR", "ANNUAL"}:
            return None, None
        return minimum, maximum
    return None, None


def _salary_payloads_from_raw(payload: Any):
    if isinstance(payload, dict):
        if payload.get("@type") == "MonetaryAmount":
            yield payload
        if isinstance(payload.get("json_ld"), dict) and payload["json_ld"].get("baseSalary"):
            yield payload["json_ld"]["baseSalary"]
        for value in payload.values():
            yield from _salary_payloads_from_raw(value)
    elif isinstance(payload, list):
        for item in payload:
            yield from _salary_payloads_from_raw(item)


def _parse_salary_payload(value: Any) -> dict[str, Any] | None:
    if isinstance(value, dict):
        return value
    text = str(value or "").strip()
    if not text or text.lower() in {"nan", "none"}:
        return None
    if "MonetaryAmount" in text:
        try:
            parsed = ast.literal_eval(text)
            return parsed if isinstance(parsed, dict) else None
        except (SyntaxError, ValueError):
            return None
    return None


def _load_json(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return value
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        return None


def _number(value: Any) -> float | None:
    if value is None:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).replace("\u202f", " ").replace(",", ".")
    match = MONEY_RE.search(text)
    if not match:
        return None
    compact = match.group(1).replace(" ", "")
    try:
        return float(compact)
    except ValueError:
        return None


def _round_salary(value: float | None) -> int | None:
    if value is None:
        return None
    return int(round(value))


def _salary_label(minimum: int | None, maximum: int | None) -> str:
    if not minimum and not maximum:
        return ""
    if minimum and maximum and minimum != maximum:
        return f"{_format_int(minimum)}-{_format_int(maximum)} €"
    return f"{_format_int(minimum or maximum)} €"


def _format_int(value: int | None) -> str:
    return "" if value is None else f"{value:,}".replace(",", " ")


def _hyperlink(url: Any, label: str | None = None) -> str:
    text = str(url or "").strip()
    if not text or text.lower() in {"nan", "none"}:
        return ""
    safe_url = text.replace('"', '""')
    safe_label = (label or text).replace('"', '""')
    return f'=HYPERLINK("{safe_url}","{safe_label}")'


def _compact_links(*values: Any) -> str:
    links: list[str] = []
    for value in values:
        for link in str(value or "").splitlines():
            link = link.strip()
            if link and link.lower() not in {"nan", "none"} and link not in links:
                links.append(link)
    return "\n".join(links[:5])


def _contract_fr(value: Any) -> str:
    text = str(value or "").strip()
    mapping = {
        "FULL_TIME": "Temps plein",
        "fulltime": "Temps plein",
        "TEMPORARY": "CDD / Temporaire",
        "temporary": "CDD / Temporaire",
        "CONTRACT": "Contrat",
        "contract": "Contrat",
        "CDI": "CDI",
        "CDD / Temporaire": "CDD / Temporaire",
    }
    parts = [mapping.get(part.strip(), part.strip()) for part in text.split(",") if part.strip()]
    deduped = []
    for part in parts:
        if part and part not in deduped:
            deduped.append(part)
    return ", ".join(deduped)


def _remote_fr(value: Any) -> str:
    if value is True or str(value).lower() == "true":
        return "Oui"
    if value is False or str(value).lower() == "false":
        return "Non"
    return ""


def _town_fr(value: Any) -> str:
    text = clean_text(value)
    if not text:
        return ""
    department_match = re.search(r"(?:^|\b)\d{2,3}\s*-\s*([^,]+)", text)
    if department_match:
        return _clean_town_name(department_match.group(1))
    parts = [part.strip() for part in text.split(",") if part.strip()]
    if len(parts) < 2:
        return _clean_town_name(text)
    trailing = " ".join(parts[1:]).casefold()
    regional_markers = (
        "france",
        "fr",
        "auvergne-rhône-alpes",
        "auvergne-rhone-alpes",
        "ara",
        "rhône-alpes",
        "rhone-alpes",
    )
    if any(marker in trailing for marker in regional_markers):
        return _clean_town_name(parts[0])
    return _clean_town_name(text)


def _clean_town_name(value: str) -> str:
    town = clean_text(value) or ""
    town = re.sub(r"\s+\d+(?:er|e|eme|ème)?\s+arrondissement\b.*$", "", town, flags=re.IGNORECASE)
    town = re.sub(r"\s+-\s+localiser\b.*$", "", town, flags=re.IGNORECASE)
    if re.match(r"^lyon(?:\b|[\s*,-])", town, flags=re.IGNORECASE):
        return "Lyon"
    return town.strip()


def _date_value(value: Any, today: date | None = None) -> date | None:
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.date()
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = str(value or "").strip()
    if not text or text.lower() in {"nat", "nan", "none"}:
        return None
    today = today or date.today()
    relative = RELATIVE_DATE_RE.search(text)
    if relative:
        unit = relative.group("unit").lower()
        count = int(relative.group("count"))
        if unit.startswith("j"):
            return today - timedelta(days=count)
        return today
    lowered = text.lower()
    if "aujourd'hui" in lowered or "aujourdhui" in lowered:
        return today
    if re.search(r"\bhier\b", lowered):
        return today - timedelta(days=1)
    match = ISO_DATE_PREFIX_RE.match(text)
    if match:
        try:
            return date.fromisoformat(match.group(0))
        except ValueError:
            return None
    match = EU_DATE_RE.match(text)
    if match:
        try:
            return date(int(match.group("year")), int(match.group("month")), int(match.group("day")))
        except ValueError:
            return None
    match = NATURAL_LANGUAGE_DATE_RE.search(text)
    if match:
        parsed = _parse_french_date_match(match, today.year)
        if parsed:
            return parsed
    return None


def _date_fr(value: Any, today: date | None = None) -> str:
    parsed = _date_value(value, today=today)
    if parsed:
        return _format_date(parsed)
    text = str(value or "").strip()
    if not text or text.lower() in {"nat", "nan", "none"}:
        return ""
    return clean_text(text) or ""


def _parse_french_date_match(match: re.Match[str], default_year: int) -> date | None:
    day = match.group("day")
    month_name = match.group("month")
    year = match.group("year") or str(default_year)
    month = FRENCH_MONTHS.get(month_name.lower())
    if not month:
        return None
    try:
        return date(int(year), month, int(day))
    except ValueError:
        return None


def _format_date(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def _excel_safe(frame: pd.DataFrame) -> pd.DataFrame:
    converted = frame.copy()
    for column in converted.columns:
        if pd.api.types.is_datetime64_any_dtype(converted[column]):
            converted[column] = converted[column].astype(str).replace({"NaT": ""})
        elif converted[column].dtype == "object":
            converted[column] = converted[column].map(_safe_cell)
    return converted


def _safe_cell(value):
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


def _column_index(frame: pd.DataFrame, column: str) -> int:
    return list(frame.columns).index(column) + 1


def _optional_column_index(frame: pd.DataFrame, column: str) -> int | None:
    if column not in frame.columns:
        return None
    return _column_index(frame, column)


def _apply_widths(worksheet, frame: pd.DataFrame, widths: dict[str, int], formats: dict[str, Any] | None = None) -> None:
    formats = formats or {}
    for index, column in enumerate(frame.columns):
        width = widths.get(column)
        if not width:
            values = [str(column), *[str(_safe_cell(value)) for value in frame[column].head(100)]]
            width = min(max(len(value) for value in values) + 2, 50)
        worksheet.set_column(index, index, width, formats.get(column))


def _autosize(worksheet, frame: pd.DataFrame) -> None:
    for index, column in enumerate(frame.columns):
        values = [str(column), *[str(_safe_cell(value)) for value in frame[column].head(100)]]
        width = min(max(len(value) for value in values) + 2, 80)
        worksheet.set_column(index, index, width)
