from __future__ import annotations

import asyncio
import os
import warnings
from collections.abc import Callable
from dataclasses import dataclass
from pathlib import Path

import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from scrapper.config import SearchConfig
from scrapper.export import export_jobs
from scrapper.scrapers import (
    scrape_aggregators,
    scrape_hellowork,
    scrape_jobspy_sources,
    scrape_web_search,
    scrape_welcome_to_the_jungle,
)
from scrapper.storage import JobStore


RUN_REAL_SMOKE_ENV = "SCRAPPER_RUN_REAL_SMOKE"
SOURCE_FILTER_ENV = "SCRAPPER_REAL_SMOKE_SOURCE"


@dataclass(frozen=True)
class SmokeSource:
    source: str
    runner: Callable[[SearchConfig, JobStore], int | dict[str, int] | None]


def _jobspy_source(source: str) -> SmokeSource:
    def run(config: SearchConfig, store: JobStore) -> dict[str, int]:
        return scrape_jobspy_sources(config.model_copy(update={"jobspy_sites": (source,)}), store)

    return SmokeSource(source, run)


def _direct_aggregator_source(source: str) -> SmokeSource:
    def run(config: SearchConfig, store: JobStore) -> dict[str, int]:
        specs = tuple(spec for spec in config.aggregators.direct_specs if spec.source == source)
        assert specs, f"No direct aggregator config found for {source}"
        aggregators = config.aggregators.model_copy(update={"direct_specs": specs, "site_search_sources": ()})
        return asyncio.run(scrape_aggregators(config.model_copy(update={"aggregators": aggregators}), store))

    return SmokeSource(source, run)


def _site_search_source(source: str) -> SmokeSource:
    def run(config: SearchConfig, store: JobStore) -> dict[str, int]:
        specs = tuple(spec for spec in config.aggregators.site_search_sources if spec.source == source)
        assert specs, f"No site-search aggregator config found for {source}"
        aggregators = config.aggregators.model_copy(update={"direct_specs": (), "site_search_sources": specs})
        return asyncio.run(scrape_aggregators(config.model_copy(update={"aggregators": aggregators}), store))

    return SmokeSource(source, run)


REAL_SMOKE_SOURCES = (
    _jobspy_source("indeed"),
    _jobspy_source("linkedin"),
    _jobspy_source("glassdoor"),
    _jobspy_source("google"),
    SmokeSource("hellowork", lambda config, store: asyncio.run(scrape_hellowork(config, store))),
    SmokeSource("welcome_to_the_jungle", lambda config, store: asyncio.run(scrape_welcome_to_the_jungle(config, store))),
    SmokeSource("web_search", lambda config, store: asyncio.run(scrape_web_search(config, store))),
    _direct_aggregator_source("example_direct_board"),
    _direct_aggregator_source("emploi_territorial"),
    _direct_aggregator_source("meteojob"),
    _direct_aggregator_source("jobijoba"),
    _direct_aggregator_source("makesense"),
    _site_search_source("apec"),
    _site_search_source("isarta"),
    _site_search_source("emploi_collectivites"),
    _site_search_source("cadremploi"),
    _site_search_source("optioncarriere"),
)


@pytest.mark.integration
@pytest.mark.slow
@pytest.mark.parametrize("smoke_source", REAL_SMOKE_SOURCES, ids=[case.source for case in REAL_SMOKE_SOURCES])
def test_real_small_scrape_by_website(smoke_source: SmokeSource, tmp_path: Path) -> None:
    if os.getenv(RUN_REAL_SMOKE_ENV) != "1":
        pytest.skip(f"Set {RUN_REAL_SMOKE_ENV}=1 to run real website-by-website scraping smoke tests.")
    selected_source = os.getenv(SOURCE_FILTER_ENV)
    if selected_source and selected_source != smoke_source.source:
        pytest.skip(f"Set {SOURCE_FILTER_ENV}={smoke_source.source} to run this website smoke test.")

    config = SearchConfig(
        queries=("example role",),
        results_wanted_per_query=1,
        custom_max_pages=1,
        custom_max_detail_pages=3,
        web_search_max_results_per_query=2,
        web_search_max_detail_pages=2,
        request_delay_secs=0.2,
        database_path=tmp_path / "job_search.duckdb",
        export_dir=tmp_path / "exports",
    )
    store = JobStore(config.database_path)
    try:
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="Unverified HTTPS request.*")
            warnings.filterwarnings("ignore", message=r"The 'text' argument to find\(\)-type methods is deprecated.*")
            warnings.filterwarnings("ignore", message="The `dict` method is deprecated; use `model_dump` instead.*")
            counts = smoke_source.runner(config, store)
        search_state = store.conn.execute(
            """
            SELECT source, last_seen_count, last_error
            FROM search_state
            ORDER BY source, query
            """
        ).fetchall()
        assert search_state, f"{smoke_source.source} did not record search state; counts={counts!r}"

        excel_path, csv_path, raw_csv_path, total = export_jobs(
            store,
            config.export_dir,
            config.export.excel_columns,
            repair_fields_enabled_by_default=config.export.repair_fields_enabled_by_default,
            repair_fields_by_source=config.export.repair_fields_by_source,
        )
        assert excel_path.exists()
        assert csv_path.exists()
        assert raw_csv_path.exists()

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["Offres"]
        assert "Données brutes" in workbook.sheetnames
        assert worksheet["A1"].value == "Postulé"
        assert worksheet["B1"].value == config.export.excel_columns[1]
        if total:
            assert "TableOffres" in worksheet.tables
            if "Statut" in config.export.excel_columns:
                status_col = config.export.excel_columns.index("Statut") + 1
                status_cell = f"{get_column_letter(status_col)}2"
                assert worksheet["A2"].value == f'={status_cell}="Postulé"'
            else:
                assert worksheet["A2"].value in {False, "☐"}

        errors = [f"{source}: {error}" for source, _seen_count, error in search_state if error]
        if errors:
            pytest.xfail("; ".join(errors))
    finally:
        store.close()
