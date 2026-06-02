from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from scrapper.config import SearchConfig
from scrapper.export import _date_fr, _date_value, _salary_bounds, export_jobs
from scrapper.location_scope import location_is_in_scope
from scrapper.models import JobOffer
from scrapper.storage import JobStore


def test_export_creates_localized_tracker_with_links_filters_and_annual_salary(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="hellowork",
                source_url="https://www.hellowork.com/fr-fr/emplois/1.html",
                direct_url="https://apply.example.com/1",
                title="Example manager",
                company="Example",
                location="Example City",
                description="Full-time example role with sample responsibilities.",
                employment_type="CDI",
                date_posted="2026-05-29T10:11:12+00:00",
                salary=str(
                    {
                        "@type": "MonetaryAmount",
                        "currency": "EUR",
                        "value": {
                            "@type": "QuantitativeValue",
                            "minValue": 3000,
                            "maxValue": 3600,
                            "unitText": "MONTH",
                        },
                    }
                ),
                raw={
                    "json_ld": {
                        "baseSalary": {
                            "@type": "MonetaryAmount",
                            "currency": "EUR",
                            "value": {
                                "@type": "QuantitativeValue",
                                "minValue": 3000,
                                "maxValue": 3600,
                                "unitText": "MONTH",
                            },
                        }
                    }
                },
            )
        )

        excel_path, csv_path, raw_csv_path, total = export_jobs(store, tmp_path / "exports")

        assert total == 1
        assert excel_path.exists()
        assert csv_path.exists()
        assert raw_csv_path.exists()

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["Offres"]
        assert workbook.sheetnames == ["Offres", "Données brutes"]
        assert list(worksheet.tables.keys()) == ["TableOffres"]
        assert worksheet.tables["TableOffres"].ref == "A1:N2"
        assert worksheet.tables["TableOffres"].autoFilter.ref == "A1:N2"
        assert worksheet.auto_filter.ref is None
        removed_headers = {
            "Statut",
            "Postuler",
            "Offre source",
            "Autres liens",
            "Source principale",
            "Doublons",
            "Description",
            "URL entreprise",
            "Email",
            "Zone recherchée",
            "Score complétude",
            "Vu le",
            "Dernière vue",
        }
        assert removed_headers.isdisjoint({cell.value for cell in worksheet[1]})
        assert worksheet["A2"].value is False
        assert worksheet["B2"].value is False
        assert worksheet["C2"].value == '=HYPERLINK("https://apply.example.com/1","https://apply.example.com/1")'
        assert worksheet["I2"].value == 36000
        assert worksheet["J2"].value == 43200
        assert worksheet["K2"].value == "36 000-43 200 €"
        published_value = worksheet["L2"].value
        assert isinstance(published_value, datetime)
        assert published_value.date() == date(2026, 5, 29)
        assert worksheet["L2"].number_format == "dd/mm/yyyy"
    finally:
        store.close()


def test_export_handles_empty_store_with_tracker_headers(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        excel_path, csv_path, raw_csv_path, total = export_jobs(store, tmp_path / "exports")

        assert total == 0
        assert excel_path.exists()
        assert csv_path.exists()
        assert raw_csv_path.exists()

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["Offres"]
        assert workbook.sheetnames == ["Offres", "Données brutes"]
        assert worksheet.max_row == 1
        assert worksheet["A1"].value == "Postulé"
        assert worksheet["B1"].value == "Rejeté"
        assert list(worksheet.tables.keys()) == []
    finally:
        store.close()


def test_export_orders_rows_by_published_date_descending_with_blanks_last(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        for title, date_posted in (
            ("Older example role", "2026-05-01"),
            ("Undated example role", None),
            ("Newer example role", "2026-06-01"),
        ):
            store.upsert_offer(
                JobOffer(
                    source="linkedin",
                    source_url=f"https://linkedin.example/jobs/{title.replace(' ', '-').lower()}",
                    title=title,
                    company="Example",
                    location="Lyon",
                    description="Full-time example role in communication.",
                    date_posted=date_posted,
                )
            )

        excel_path, _, _, total = export_jobs(store, tmp_path / "exports")

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["Offres"]
        titles = [worksheet[f"D{row}"].value for row in range(2, 5)]

        assert total == 3
        assert titles == ["Newer example role", "Older example role", "Undated example role"]
    finally:
        store.close()


def test_export_uses_configured_excel_columns(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="hellowork",
                source_url="https://www.hellowork.com/fr-fr/emplois/1.html",
                direct_url="https://apply.example.com/1",
                title="Example manager",
                company="Example",
                location="Example City",
                description="Full-time example role with sample responsibilities.",
            )
        )

        excel_path, csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Intitulé", "Entreprise", "Lien"),
        )

        workbook = load_workbook(excel_path, data_only=False)
        worksheet = workbook["Offres"]
        assert worksheet.tables["TableOffres"].ref == "A1:C2"
        assert worksheet.tables["TableOffres"].autoFilter.ref == "A1:C2"
        assert [cell.value for cell in worksheet[1]] == ["Intitulé", "Entreprise", "Lien"]
        assert worksheet["A2"].value == "Example manager"
        assert worksheet["B2"].value == "Example"
        assert worksheet["C2"].value == '=HYPERLINK("https://apply.example.com/1","https://apply.example.com/1")'
        assert worksheet.auto_filter.ref is None
        assert pd.read_csv(csv_path).columns.tolist() == ["Intitulé", "Entreprise", "Lien"]
    finally:
        store.close()


def test_export_adds_status_dropdown_validation(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="hellowork",
                source_url="https://www.hellowork.com/fr-fr/emplois/1.html",
                title="Example manager",
                company="Example",
                location="Example City",
            )
        )

        excel_path, _csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Statut", "Lien", "Intitulé", "Entreprise"),
        )

        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        validations = list(worksheet.data_validations.dataValidation)
        assert len(validations) == 1
        assert validations[0].type == "list"
        assert validations[0].formula1 == '"En Attente,Postulé,Refusé"'
        assert "A2:A1048576" in str(validations[0].sqref)
        assert worksheet["A2"].value == "En Attente"
    finally:
        store.close()


def test_export_preserves_existing_status_for_matching_tracker_rows(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/1",
                title="Example manager",
                company="Example",
                location="Lyon",
            )
        )
        excel_path, _csv_path, _raw_csv_path, _total = export_jobs(
            store,
            export_dir,
            excel_columns=("Statut", "Lien", "Intitulé", "Entreprise"),
        )
    finally:
        store.close()

    workbook = load_workbook(excel_path, data_only=False)
    worksheet = workbook["Offres"]
    worksheet["A2"] = "Postulé"
    workbook.save(excel_path)
    workbook.close()

    refreshed_store = JobStore(tmp_path / "refreshed.duckdb")
    try:
        refreshed_store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/1?utm_source=refresh",
                title="Example manager",
                company="Example",
                location="Lyon",
            )
        )

        excel_path, csv_path, _raw_csv_path, total = export_jobs(
            refreshed_store,
            export_dir,
            excel_columns=("Statut", "Lien", "Intitulé", "Entreprise"),
        )

        assert total == 1
        assert pd.read_csv(csv_path).loc[0, "Statut"] == "Postulé"
        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet["A2"].value == "Postulé"
    finally:
        refreshed_store.close()


def test_export_retains_existing_tracker_rows_when_current_export_has_new_rows(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    old_store = JobStore(tmp_path / "old.duckdb")
    try:
        old_store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/old",
                title="Old example manager",
                company="Example",
                location="Lyon",
            )
        )
        export_jobs(old_store, export_dir, excel_columns=("Statut", "Lien", "Intitulé", "Entreprise"))
    finally:
        old_store.close()

    new_store = JobStore(tmp_path / "new.duckdb")
    try:
        new_store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/new",
                title="New example manager",
                company="Example",
                location="Lyon",
            )
        )

        excel_path, csv_path, _raw_csv_path, total = export_jobs(
            new_store,
            export_dir,
            excel_columns=("Statut", "Lien", "Intitulé", "Entreprise"),
        )

        assert total == 2
        frame = pd.read_csv(csv_path)
        assert frame["Intitulé"].to_list() == ["New example manager", "Old example manager"]
        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet.tables["TableOffres"].ref == "A1:D3"
    finally:
        new_store.close()


def test_export_rewrites_retained_ouvrir_links_to_full_url_labels(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    old_store = JobStore(tmp_path / "old.duckdb")
    try:
        old_store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/old",
                title="Old example manager",
                company="Example",
                location="Lyon",
            )
        )
        excel_path, _csv_path, _raw_csv_path, _total = export_jobs(old_store, export_dir)
    finally:
        old_store.close()

    workbook = load_workbook(excel_path, data_only=False)
    worksheet = workbook["Offres"]
    worksheet["C2"] = '=HYPERLINK("https://linkedin.example/jobs/old","Ouvrir")'
    worksheet["A2"] = '=B2="Postulé"'
    workbook.save(excel_path)
    workbook.close()

    empty_store = JobStore(tmp_path / "empty.duckdb")
    try:
        excel_path, csv_path, _raw_csv_path, total = export_jobs(empty_store, export_dir)

        assert total == 1
        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet["A2"].value is False
        assert worksheet["C2"].value == (
            '=HYPERLINK("https://linkedin.example/jobs/old","https://linkedin.example/jobs/old")'
        )
        assert "Ouvrir" not in pd.read_csv(csv_path).loc[0, "Lien"]
    finally:
        empty_store.close()


def test_export_preserves_existing_checkbox_for_matching_tracker_rows(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/1",
                title="Example manager",
                company="Example",
                location="Lyon",
            )
        )
        excel_path, _csv_path, _raw_csv_path, _total = export_jobs(store, export_dir)
    finally:
        store.close()

    workbook = load_workbook(excel_path, data_only=False)
    worksheet = workbook["Offres"]
    worksheet["A2"] = True
    workbook.save(excel_path)
    workbook.close()

    refreshed_store = JobStore(tmp_path / "refreshed.duckdb")
    try:
        refreshed_store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://linkedin.example/jobs/1?utm_source=refresh",
                title="Example manager",
                company="Example",
                location="Lyon",
            )
        )

        excel_path, csv_path, _raw_csv_path, total = export_jobs(refreshed_store, export_dir)

        assert total == 1
        assert bool(pd.read_csv(csv_path).loc[0, "Postulé"])
        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet["A2"].value is True
        assert worksheet["C2"].value == '=HYPERLINK("https://linkedin.example/jobs/1?utm_source=refresh","https://linkedin.example/jobs/1?utm_source=refresh")'
    finally:
        refreshed_store.close()


def test_export_does_not_retain_existing_rows_rejected_by_current_filters(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    old_store = JobStore(tmp_path / "old.duckdb")
    try:
        old_store.upsert_offer(
            JobOffer(
                source="indeed",
                source_url="https://indeed.example/jobs/alternance",
                title="Offre d’alternance / Community Manager (H/F)",
                company="Example",
                employment_type="Temps plein",
            )
        )
        export_jobs(old_store, export_dir)
    finally:
        old_store.close()

    empty_store = JobStore(tmp_path / "empty.duckdb")
    try:
        excel_path, csv_path, _raw_csv_path, total = export_jobs(
            empty_store,
            export_dir,
            include_terms=("community",),
            exclude_terms=("alternance",),
        )

        assert total == 0
        assert pd.read_csv(csv_path).empty
        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet.max_row == 1
        assert list(worksheet.tables.keys()) == []
    finally:
        empty_store.close()


def test_export_does_not_retain_existing_rows_rejected_by_location_scope(tmp_path: Path) -> None:
    export_dir = tmp_path / "exports"
    config = SearchConfig()
    old_store = JobStore(tmp_path / "old.duckdb")
    try:
        old_store.upsert_offer(
            JobOffer(
                source="indeed",
                source_url="https://indeed.example/jobs/thiers",
                title="Example manager",
                company="Example",
                location="Thiers, Auvergne-Rhône-Alpes, France",
            )
        )
        export_jobs(old_store, export_dir)
    finally:
        old_store.close()

    empty_store = JobStore(tmp_path / "empty.duckdb")
    try:
        _excel_path, csv_path, _raw_csv_path, total = export_jobs(
            empty_store,
            export_dir,
            location_allowed=lambda location: location_is_in_scope(location, config.location_scope),
        )

        assert total == 0
        assert pd.read_csv(csv_path).empty
    finally:
        empty_store.close()


def test_export_keeps_postule_boolean_when_status_column_is_omitted(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="hellowork",
                source_url="https://www.hellowork.com/fr-fr/emplois/1.html",
                title="Example manager",
                company="Example",
                location="Example City",
                description="Full-time example role with sample responsibilities.",
            )
        )

        excel_path, _csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Postulé", "Lien", "Intitulé", "Entreprise"),
        )

        worksheet = load_workbook(excel_path, data_only=False)["Offres"]
        assert worksheet.tables["TableOffres"].ref == "A1:D2"
        assert worksheet.tables["TableOffres"].autoFilter.ref == "A1:D2"
        assert [cell.value for cell in worksheet[1]] == ["Postulé", "Lien", "Intitulé", "Entreprise"]
        assert worksheet["A2"].value is False
        assert worksheet["B2"].value == (
            '=HYPERLINK("https://www.hellowork.com/fr-fr/emplois/1.html",'
            '"https://www.hellowork.com/fr-fr/emplois/1.html")'
        )
        assert worksheet.auto_filter.ref is None
    finally:
        store.close()


def test_export_sanitizes_tracker_fields_for_france_travail_style_rows(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="france_travail",
                source_url="https://candidat.francetravail.fr/offres/recherche/detail/208WKXK",
                direct_url="https://candidat.francetravail.fr/offres/recherche/detail/208WKXK",
                title="208WKXK Chargé de communication (F/H)",
                company="69 - Pusignan - Localiser avec Mappy",
                location="Lyon",
                description="Votre objectif principal sera de mettre en œuvre la stratégie de communication.",
            )
        )

        _excel_path, csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Intitulé", "Entreprise", "Lieu"),
        )

        frame = pd.read_csv(csv_path).fillna("")
        assert frame.iloc[0].to_dict() == {
            "Intitulé": "Chargé de communication (F/H)",
            "Entreprise": "",
            "Lieu": "Pusignan",
        }
    finally:
        store.close()


def test_export_can_disable_repairs_for_specific_source(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="france_travail",
                source_url="https://candidat.francetravail.fr/offres/recherche/detail/208WKXK",
                direct_url="https://candidat.francetravail.fr/offres/recherche/detail/208WKXK",
                title="208WKXK Chargé de communication (F/H)",
                company="69 - Pusignan - Localiser avec Mappy",
                location="Lyon",
                description="Votre objectif principal sera de mettre en œuvre la stratégie de communication.",
            )
        )

        _excel_path, csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Intitulé", "Entreprise", "Lieu"),
            repair_fields_enabled_by_default=True,
            repair_fields_by_source={"france_travail": False},
        )

        frame = pd.read_csv(csv_path).fillna("")
        assert frame.iloc[0].to_dict() == {
            "Intitulé": "208WKXK Chargé de communication (F/H)",
            "Entreprise": "69 - Pusignan - Localiser avec Mappy",
            "Lieu": "Lyon",
        }
    finally:
        store.close()


def test_export_normalizes_location_to_town_for_tracker_display(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://www.linkedin.com/jobs/view/1",
                title="Chargé communication",
                company="Example",
                location="Lyon, Auvergne-Rhône-Alpes, France",
            )
        )
        store.upsert_offer(
            JobOffer(
                source="linkedin",
                source_url="https://www.linkedin.com/jobs/view/2",
                title="Responsable communication",
                company="Example",
                location="Montbonnot-Saint-Martin, ARA, FR",
            )
        )
        store.upsert_offer(
            JobOffer(
                source="france_travail",
                source_url="https://candidat.francetravail.fr/offres/recherche/detail/1",
                title="Community manager",
                company="Example",
                location="Chargé(e) de Communication - Relations Presse (H/F) 69 - Lyon 2e Arrondissement",
            )
        )
        store.upsert_offer(
            JobOffer(
                source="jobijoba",
                source_url="https://www.jobijoba.com/fr/annonce/54/1",
                title="Content manager",
                company="Example",
                location="Lyon 9e",
            )
        )

        _excel_path, csv_path, _raw_csv_path, _total = export_jobs(
            store,
            tmp_path / "exports",
            excel_columns=("Intitulé", "Entreprise", "Lieu"),
        )

        frame = pd.read_csv(csv_path).sort_values("Intitulé").reset_index(drop=True)
        assert frame["Lieu"].tolist() == ["Lyon", "Lyon", "Lyon", "Montbonnot-Saint-Martin"]
    finally:
        store.close()


def test_export_rejects_unknown_excel_columns(tmp_path: Path) -> None:
    store = JobStore(tmp_path / "jobs.duckdb")
    try:
        with pytest.raises(ValueError, match="Unknown Excel column"):
            export_jobs(store, tmp_path / "exports", excel_columns=("Intitulé", "Bad column"))
    finally:
        store.close()


def test_salary_bounds_leave_untrusted_non_annual_units_empty() -> None:
    row = pd.Series(
        {
            "min_amount": None,
            "max_amount": None,
            "salary": {
                "@type": "MonetaryAmount",
                "currency": "EUR",
                "value": {
                    "@type": "QuantitativeValue",
                    "minValue": 20,
                    "maxValue": 30,
                    "unitText": "HOUR",
                },
            },
            "raw_json": None,
        }
    )

    assert _salary_bounds(row) == (None, None)


def test_date_fr_formats_iso_natural_language_and_relative_dates() -> None:
    today = date(2026, 6, 4)

    assert _date_fr("2026-05-29T10:11:12+00:00", today=today) == "29/05/2026"
    assert _date_fr("Publié le 28 mai 2026", today=today) == "28/05/2026"
    assert _date_fr("29 mai Description de l'offre", today=today) == "29/05/2026"
    assert _date_fr("24 avril 2026 Le Syndic", today=today) == "24/04/2026"
    assert _date_fr("Il y a 10 h Description", today=today) == "04/06/2026"
    assert _date_fr("Il y a 2 jours", today=today) == "02/06/2026"
    assert _date_value("29/05/2026", today=today) == date(2026, 5, 29)
